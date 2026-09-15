"""Document text extraction: PDF / DOCX / TXT / code / CSV / JSON / PPTX / RTF with clear errors."""
from __future__ import annotations
import io
import json as _json

MAX_CHARS = 200_000

# Extensions treated as plain UTF-8 text
_TEXT_EXTS = frozenset({
    "txt", "md", "csv", "json", "py", "java", "js", "ts", "jsx", "tsx",
    "sql", "html", "css", "scss", "less", "rb", "go", "rs", "cpp", "c",
    "h", "hpp", "sh", "bash", "yaml", "yml", "toml", "ini", "cfg", "conf",
    "xml", "svg", "php", "swift", "kt", "kts", "dart", "r", "lua", "pl",
    "ex", "exs", "hs", "ml", "clj", "lisp", "vim",
})


def extract_text(filename: str, data: bytes) -> str:
    ext = (filename.rsplit(".", 1)[-1] if "." in filename else "").lower()
    if ext == "pdf":
        return _pdf(data, filename)
    if ext == "docx":
        return _docx(data, filename)
    if ext == "pptx":
        return _pptx(data, filename)
    if ext == "rtf":
        return _rtf(data, filename)
    if ext in _TEXT_EXTS:
        return _text(data, filename)
    raise ValueError(
        f"Unsupported document type '.{ext}' for {filename}. "
        "Use PDF, DOCX, PPTX, TXT, MD, CSV, JSON, or a code file."
    )


def _cap(text: str, filename: str) -> str:
    text = (text or "").strip()
    if not text:
        raise ValueError(f"{filename} appears empty or is a scanned image PDF without text.")
    if len(text) > MAX_CHARS:
        text = text[:MAX_CHARS] + "\n\n[TRUNCATED: document too large, first 200k chars kept]"
    return text


def _text(data: bytes, filename: str) -> str:
    try:
        text = data.decode("utf-8", errors="replace")
    except Exception as exc:
        raise ValueError(f"Could not read {filename}: {exc}") from exc
    return _cap(text, filename)


def _pdf(data: bytes, filename: str) -> str:
    try:
        from pypdf import PdfReader
    except Exception as exc:
        raise ValueError(f"PDF support unavailable: {exc}") from exc
    try:
        reader = PdfReader(io.BytesIO(data))
        parts = [(p.extract_text() or "") for p in reader.pages]
    except Exception as exc:
        raise ValueError(f"Corrupted or unreadable PDF {filename}: {exc}") from exc
    return _cap("\n".join(parts), filename)


def _docx(data: bytes, filename: str) -> str:
    try:
        import docx
    except Exception as exc:
        raise ValueError(f"DOCX support unavailable: {exc}") from exc
    try:
        doc = docx.Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
    except Exception as exc:
        raise ValueError(f"Corrupted or unreadable DOCX {filename}: {exc}") from exc
    return _cap("\n".join(parts), filename)


def _pptx(data: bytes, filename: str) -> str:
    try:
        from pptx import Presentation
    except Exception as exc:
        raise ValueError(f"PPTX support unavailable: {exc}") from exc
    try:
        prs = Presentation(io.BytesIO(data))
        parts = []
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    parts.append(shape.text)
    except Exception as exc:
        raise ValueError(f"Corrupted or unreadable PPTX {filename}: {exc}") from exc
    return _cap("\n".join(parts), filename)


def _rtf(data: bytes, filename: str) -> str:
    """Strip RTF control words and extract plain text."""
    try:
        raw = data.decode("utf-8", errors="replace")
    except Exception as exc:
        raise ValueError(f"Could not read {filename}: {exc}") from exc
    import re
    # Remove RTF header group and control words
    text = re.sub(r"\\[a-z]+\d*\s?", "", raw)
    text = re.sub(r"\{[^}]*\}", "", text)
    text = re.sub(r"\\[^a-z]", "", text)
    text = re.sub(r"[ ]+", " ", text)
    return _cap(text.strip(), filename)


def chunk(text: str, size: int = 4000, overlap: int = 400) -> list[str]:
    chunks, i = [], 0
    while i < len(text):
        chunks.append(text[i:i + size])
        i += size - overlap
    return chunks
